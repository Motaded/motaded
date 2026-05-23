#!/usr/bin/env python3
"""Build article_import_official.catalog.php from the import spreadsheet."""

from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Install openpyxl: pip install openpyxl") from exc

ROW_ID = 0
ROW_LANG = 1
ROW_PUBLISHED = 3
ROW_TITLE = 4
ROW_ALIAS = 5
ROW_BODY = 6
ROW_IMAGE = 7

PAIR_RE = re.compile(r"^(\d+)_(en|ar)$")


def php_export(value) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    text = str(value)
    text = text.replace("\\", "\\\\").replace("'", "\\'")
    return f"'{text}'"


def main() -> int:
    root = Path(__file__).resolve().parents[5]
    xlsx = root / "FINAL_IMPORT_CLEAN_V2.xlsx"
    if len(sys.argv) > 1:
        xlsx = Path(sys.argv[1]).expanduser().resolve()
    if not xlsx.is_file():
        print(f"Spreadsheet not found: {xlsx}", file=sys.stderr)
        return 1

    out = (
        Path(__file__).resolve().parents[1]
        / "data"
        / "article_import_official.catalog.php"
    )

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if "Import Ready" not in wb.sheetnames:
        print("Sheet 'Import Ready' not found.", file=sys.stderr)
        return 1
    ws = wb["Import Ready"]

    pairs: dict[str, dict[str, dict]] = {}
    skipped = 0
    truncated: list[str] = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        raw_id = row[ROW_ID]
        if raw_id is None:
            continue
        row_id = str(raw_id).strip()
        match = PAIR_RE.match(row_id)
        if not match:
            skipped += 1
            continue
        pair_key, lang = match.group(1), match.group(2)
        body = row[ROW_BODY] if row[ROW_BODY] is not None else ""
        body = str(body)
        if len(body) >= 32767:
            truncated.append(row_id)

        entry = {
            "title": str(row[ROW_TITLE] or "").strip(),
            "path_alias": str(row[ROW_ALIAS] or "").strip(),
            "body": body,
            "published": int(row[ROW_PUBLISHED] or 0),
            "image": str(row[ROW_IMAGE] or "").strip(),
        }
        pairs.setdefault(pair_key, {})[lang] = entry

    wb.close()

    incomplete = [k for k, v in pairs.items() if "en" not in v or "ar" not in v]
    if incomplete:
        print(f"Incomplete pairs (missing en/ar): {', '.join(incomplete[:20])}", file=sys.stderr)
        return 1

    lines = [
        "<?php",
        "",
        "/**",
        " * @file",
        " * Article import catalog generated from spreadsheet.",
        " *",
        " * Regenerate:",
        " *   python3 web/modules/custom/motaded_custom/scripts/build_article_import_catalog.py",
        " */",
        "",
        "declare(strict_types=1);",
        "",
        "/**",
        " * @return array<string, array{en: array<string, mixed>, ar: array<string, mixed>}>",
        " */",
        "return [",
    ]

    for key in sorted(pairs, key=lambda k: int(k)):
        lines.append(f"  '{key}' => [")
        for lang in ("en", "ar"):
            row = pairs[key][lang]
            lines.append(f"    '{lang}' => [")
            for field, val in row.items():
                lines.append(f"      '{field}' => {php_export(val)},")
            lines.append("    ],")
        lines.append("  ],")

    lines.append("];")
    lines.append("")

    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(lines), encoding="utf-8")

    print(f"Wrote {len(pairs)} article pair(s) to {out}")
    print(f"Skipped {skipped} non-article row(s).")
    if truncated:
        print(
            f"WARNING: {len(truncated)} row(s) at Excel 32,767 char limit (body may be truncated): "
            + ", ".join(truncated[:10])
            + ("…" if len(truncated) > 10 else "")
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
